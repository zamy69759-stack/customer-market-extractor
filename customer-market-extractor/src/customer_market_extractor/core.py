from __future__ import annotations

import csv
import email
import hashlib
import html
import json
import re
from dataclasses import dataclass, asdict
from datetime import date
from email import policy
from pathlib import Path

TEXT_EXTS = {".txt", ".md", ".html", ".htm", ".srt", ".vtt", ".json"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}
AUDIO_EXTS = {".mp3", ".wav", ".m4a"}
OPTIONAL_EXTS = {".docx", ".xlsx", ".pdf", ".pptx", ".msg"}


@dataclass
class ParseResult:
    path: str
    sha256: str
    kind: str
    status: str
    reason: str
    text: str = ""


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def safe_relative(root: Path, requested: str) -> Path:
    root = root.resolve()
    target = (root / requested).resolve()
    if target != root and root not in target.parents:
        raise ValueError("请求路径越出已配置的原始语料库")
    return target


def discover(root: Path) -> list[Path]:
    if not root.exists():
        raise FileNotFoundError(f"路径不可用: {root}")
    if not root.is_dir():
        return [root]
    return sorted(
        p for p in root.rglob("*")
        if p.is_file() and not p.name.startswith("._") and p.name not in {".DS_Store", "Thumbs.db"}
    )


def _strip_html(value: str) -> str:
    value = re.sub(r"(?is)<(script|style).*?>.*?</\\1>", "", value)
    return html.unescape(re.sub(r"(?s)<[^>]+>", " ", value))


def parse_file(path: Path, display_path: str | None = None) -> ParseResult:
    digest = sha256(path)
    ext = path.suffix.lower()
    shown = display_path or str(path)
    try:
        if ext == ".eml":
            with path.open("rb") as stream:
                msg = email.message_from_binary_file(stream, policy=policy.default)
            headers = [f"{k}: {msg.get(k, '')}" for k in ("Date", "From", "To", "Subject", "Message-ID", "In-Reply-To")]
            parts = []
            for part in msg.walk():
                if part.get_content_type() == "text/plain" and not part.get_filename():
                    parts.append(part.get_content())
            return ParseResult(shown, digest, "email", "成功", "保留邮件头与纯文本正文", "\n".join(headers + parts))
        if ext == ".csv":
            with path.open(encoding="utf-8-sig", newline="") as f:
                rows = list(csv.reader(f))
            return ParseResult(shown, digest, "chat/table", "成功", "CSV 行号可定位", "\n".join(" | ".join(r) for r in rows))
        if ext in TEXT_EXTS:
            raw = path.read_text(encoding="utf-8-sig", errors="replace")
            text = _strip_html(raw) if ext in {".html", ".htm"} else raw
            return ParseResult(shown, digest, "text/transcript", "成功", "按行号定位", text)
        if ext in IMAGE_EXTS:
            return ParseResult(shown, digest, "image", "部分成功", "已登记图片；默认未获授权进行 OCR")
        if ext in AUDIO_EXTS:
            return ParseResult(shown, digest, "audio", "部分成功", "已登记音频；需明确授权且检测到转写能力")
        if ext in OPTIONAL_EXTS:
            return _parse_optional(path, shown, digest, ext)
        return ParseResult(shown, digest, "unknown", "失败", f"不支持的文件类型: {ext or '无扩展名'}")
    except Exception as exc:
        return ParseResult(shown, digest, ext.lstrip(".") or "unknown", "失败", f"解析错误: {exc}")


def _parse_optional(path: Path, shown: str, digest: str, ext: str) -> ParseResult:
    if ext == ".docx":
        from docx import Document
        text = "\n".join(p.text for p in Document(path).paragraphs)
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        text = "\n".join(f"[{ws.title}!{row[0].row}] " + " | ".join(str(c.value or "") for c in row) for ws in wb for row in ws.iter_rows())
    elif ext == ".pdf":
        from pypdf import PdfReader
        pages = [f"[page {i}] {p.extract_text() or ''}" for i, p in enumerate(PdfReader(path).pages, 1)]
        text = "\n".join(pages)
        if not text.strip():
            return ParseResult(shown, digest, "pdf", "部分成功", "PDF 无可提取文字层，需要 OCR")
    elif ext == ".pptx":
        from pptx import Presentation
        text = "\n".join(f"[slide {i}] " + " ".join(s.text for s in slide.shapes if hasattr(s, "text")) for i, slide in enumerate(Presentation(path).slides, 1))
    elif ext == ".msg":
        import extract_msg
        msg = extract_msg.Message(path)
        text = f"Date: {msg.date}\nFrom: {msg.sender}\nTo: {msg.to}\nSubject: {msg.subject}\n\n{msg.body or ''}"
    else:
        raise RuntimeError("unsupported optional parser")
    return ParseResult(shown, digest, ext.lstrip("."), "成功", "本地解析器可用", text)


def load_config(skill_dir: Path) -> dict:
    config_path = skill_dir / "config.json"
    example_path = skill_dir / "config.example.json"
    if config_path.exists():
        return json.loads(config_path.read_text(encoding="utf-8"))
    if example_path.exists():
        return json.loads(example_path.read_text(encoding="utf-8"))
    return {
        "source_root": None,
        "knowledge_root": None,
        "work_root": None,
        "knowledge_category": "客户与市场",
        "created_at": None,
    }


def sanitize(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")
    return value or "未命名"


SECTIONS = [
    "01 案例一句话概括", "02 客户与订单背景", "03 客户提出的问题",
    "04 业务员当时的想法与难点（仅有证据时填写）", "05 客户真实意图分析（推断须显式标注）",
    "06 处理思路", "07 关键沟通过程", "08 实际执行与客户反馈", "09 当前结果与未确认事项",
    "10 可复用客户与市场知识", "11 LinkedIn 内容原材料", "12 原始证据与知识库检索问答",
]


def card_template(topic: str, obj: str, source_files: list[str], event_id: str, created: date | None = None) -> str:
    created = created or date.today()
    sources = json.dumps(source_files, ensure_ascii=False)
    front = f'''---\n知识类型: 客户与市场\n知识主题: "{topic}"\n对象名称: "{obj}"\n首次建档日期: {created.isoformat()}\n所属市场: ""\n关联产品: []\n内容类型: "真实业务案例"\n案例完整度: "fragment"\n来源类型: []\n来源文件: {sources}\n关联事件ID: "{event_id}"\n事实证据等级: "direct"\n审核状态: 待审核\n脱敏状态: 待审核\n公开发布许可: 待确认\n关键词: []\n---\n\n# {topic}\n'''
    return front + "\n".join(f"\n## {s}\n\n原始资料未提供\n" for s in SECTIONS)


def validate_card(path: Path) -> list[str]:
    errors = []
    if not re.fullmatch(r"MKT_[^_]+_[^_]+_\d{8}\.md", path.name):
        errors.append("文件名不符合 MKT 四段式")
    text = path.read_text(encoding="utf-8")
    for key in ["知识类型", "知识主题", "对象名称", "首次建档日期", "关联事件ID", "审核状态", "公开发布许可"]:
        if f"{key}:" not in text:
            errors.append(f"缺少 YAML 键: {key}")
    for section in SECTIONS:
        if f"## {section}" not in text:
            errors.append(f"缺少章节: {section}")
    return errors


def result_dict(result: ParseResult) -> dict:
    d = asdict(result)
    d.pop("text", None)
    return d
